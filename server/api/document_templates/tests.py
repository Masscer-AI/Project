import json
import os
import tempfile
import uuid

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from rest_framework.test import APIClient

from api.ai_layers.models import Agent, LanguageModel
from api.authenticate.models import Organization, Token, UserProfile
from api.document_templates.models import DocumentTemplate
from api.document_templates.rendering import render_docx_template_to_bytes
from api.document_templates.utils import (
    build_template_metadata,
    extract_placeholders_from_docx_path,
    merge_variables_metadata,
)
from api.messaging.models import Conversation
from api.providers.models import AIProvider
from docx import Document


def _minimal_docx_bytes_with_placeholder(placeholder: str = "client_name") -> bytes:
    doc = Document()
    doc.add_paragraph(f"Hello {{{{ {placeholder} }}}} there.")
    path = tempfile.mktemp(suffix=".docx")
    doc.save(path)
    try:
        with open(path, "rb") as f:
            return f.read()
    finally:
        os.unlink(path)


@override_settings(MEDIA_ROOT=tempfile.mkdtemp())
class DocumentTemplateUtilsTests(TestCase):
    def test_extract_and_merge_metadata(self):
        raw = _minimal_docx_bytes_with_placeholder("invoice_id")
        path = tempfile.mktemp(suffix=".docx")
        with open(path, "wb") as f:
            f.write(raw)
        try:
            ph = extract_placeholders_from_docx_path(path)
            self.assertIn("invoice_id", ph)
            prev = {
                "placeholders": ["invoice_id"],
                "variables": {
                    "invoice_id": {"description": "keep me", "required": False, "example": "INV-1"}
                },
            }
            md = build_template_metadata(ph, prev)
            self.assertEqual(md["placeholders"], ["invoice_id"])
            self.assertEqual(md["variables"]["invoice_id"]["description"], "keep me")
        finally:
            os.unlink(path)

    def test_merge_variables_drops_stale(self):
        prev_vars = {"old": {"description": "gone", "required": True, "example": ""}}
        merged = merge_variables_metadata(prev_vars, ["new_key"])
        self.assertIn("new_key", merged)
        self.assertNotIn("old", merged)


@override_settings(MEDIA_ROOT=tempfile.mkdtemp())
class DocumentTemplateAPITests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.user = User.objects.create_user(
            username="tpl-owner",
            email="tpl-owner@example.com",
            password="pass-123456",
        )
        self.org = Organization.objects.create(name="Org TPL", owner=self.user)
        UserProfile.objects.create(user=self.user, organization=self.org, name="Owner")
        self.token = Token.objects.create(user=self.user, token_type="permanent")
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {self.token.key}")

        provider = AIProvider.objects.create(name="OpenAI")
        self.llm = LanguageModel.objects.create(
            provider=provider,
            slug="gpt-4o-mini",
            name="GPT 4o mini",
        )
        self.agent = Agent.objects.create(
            name="Org Agent",
            salute="hi",
            act_as="assistant",
            organization=self.org,
            llm=self.llm,
            model_slug=self.llm.slug,
            model_provider="openai",
        )
        self.conv = Conversation.objects.create(
            user=self.user,
            organization=self.org,
            title="C1",
        )

    def test_upload_list_patch_variables_forbidden_other_user(self):
        docx = _minimal_docx_bytes_with_placeholder("field_a")
        up = SimpleUploadedFile(
            "t.docx",
            docx,
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        url = f"/v1/document-templates/organizations/{self.org.id}/templates/"
        r = self.client.post(
            url,
            data={"name": "T1", "description": "d", "file": up},
            format="multipart",
        )
        self.assertEqual(r.status_code, 201, r.content)
        tid = r.json()["template"]["id"]
        self.assertIn("field_a", r.json()["template"]["metadata"]["placeholders"])

        r2 = self.client.get(url)
        self.assertEqual(r2.status_code, 200)
        self.assertEqual(len(r2.json()["templates"]), 1)

        vr = self.client.patch(
            f"{url}{tid}/variables/",
            data=json.dumps(
                {
                    "variables": {
                        "field_a": {"description": "Field A desc", "required": True},
                    }
                }
            ),
            content_type="application/json",
        )
        self.assertEqual(vr.status_code, 200)
        self.assertEqual(
            vr.json()["template"]["metadata"]["variables"]["field_a"]["description"],
            "Field A desc",
        )

        other = User.objects.create_user(username="other", email="o@e.com", password="x")
        tok2 = Token.objects.create(user=other, token_type="permanent")
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {tok2.key}")
        r403 = self.client.get(url)
        self.assertEqual(r403.status_code, 403)

    def test_assignment_and_render_creates_attachment(self):
        docx = _minimal_docx_bytes_with_placeholder("client_name")
        up = SimpleUploadedFile(
            "c.docx",
            docx,
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        base = f"/v1/document-templates/organizations/{self.org.id}/templates/"
        r = self.client.post(
            base,
            data={"name": "Contract", "description": "", "file": up},
            format="multipart",
        )
        self.assertEqual(r.status_code, 201)
        tid = r.json()["template"]["id"]

        self.client.patch(
            f"{base}{tid}/variables/",
            data=json.dumps(
                {"variables": {"client_name": {"description": "Legal name", "required": True}}}
            ),
            content_type="application/json",
        )

        aurl = f"/v1/document-templates/agents/{self.agent.slug}/template-assignments/"
        ar = self.client.post(
            aurl,
            data=json.dumps({"template_id": tid, "usage_instructions": "When user asks for contract"}),
            content_type="application/json",
        )
        self.assertEqual(ar.status_code, 201, ar.content)

        tpl = DocumentTemplate.objects.get(id=tid)
        out = render_docx_template_to_bytes(tpl, {"client_name": "ACME Ltd"})
        self.assertTrue(out.startswith(b"PK"))

        from api.ai_layers.tools.render_document_template import _render_impl

        res = _render_impl(
            template_id=str(tid),
            variables={"client_name": "ACME Ltd"},
            output_filename="out.docx",
            conversation_id=str(self.conv.id),
            user_id=self.user.id,
            agent_slug=self.agent.slug,
            organization_id=str(self.org.id),
        )
        self.assertTrue(res.attachment_id)
        from api.messaging.models import MessageAttachment

        att = MessageAttachment.objects.get(id=uuid.UUID(res.attachment_id))
        self.assertEqual(att.conversation_id, self.conv.id)
        self.assertTrue(att.file)

    def test_assignment_rejects_cross_org_template(self):
        other_org = Organization.objects.create(name="O2", owner=self.user)
        tpl = DocumentTemplate.objects.create(
            organization=other_org,
            created_by=self.user,
            name="Foreign",
            description="",
            metadata={
                "placeholders": ["x"],
                "variables": {"x": {"description": "", "required": True, "example": ""}},
            },
        )
        tpl.file.save(
            "f.docx",
            SimpleUploadedFile(
                "f.docx",
                b"PK\x03\x04" + b"\x00" * 200,
                content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ),
        )
        tpl.save()

        aurl = f"/v1/document-templates/agents/{self.agent.slug}/template-assignments/"
        ar = self.client.post(
            aurl,
            data=json.dumps({"template_id": str(tpl.id)}),
            content_type="application/json",
        )
        self.assertEqual(ar.status_code, 404)
