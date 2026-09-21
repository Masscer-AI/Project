from __future__ import annotations

import csv
import io
from typing import Any

from api.org_lists.schemas import ColumnConfig, OrganizationListConfig, validate_list_config_for_storage
from api.utils.spreadsheet_tools import OLE_COMPOUND_MAGIC

# Safety cap on imported *data* rows after the header (SAT 69-B can exceed 10k).
_MAX_DATA_ROWS = 500_000
_MAX_COLS = 256
_HEADER_SCAN_ROWS = 30
_MAX_COLLECTED_ROWS = _HEADER_SCAN_ROWS + _MAX_DATA_ROWS

CSV_EXTENSIONS = {".csv", ".txt"}
EXCEL_XLSX_EXTENSIONS = {".xlsx"}
EXCEL_XLS_EXTENSIONS = {".xls"}


class TabularParseError(ValueError):
    pass


def _normalize_cell(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _detect_format(raw: bytes, filename: str, content_type: str) -> str:
    name = (filename or "").lower()
    ext = ""
    if "." in name:
        ext = name[name.rfind(".") :]
    ctype = (content_type or "").lower().split(";")[0].strip()

    if ext in CSV_EXTENSIONS or ctype in ("text/csv", "application/csv"):
        return "csv"
    if ext in EXCEL_XLSX_EXTENSIONS or "spreadsheetml" in ctype:
        return "xlsx"
    if ext in EXCEL_XLS_EXTENSIONS or ctype == "application/vnd.ms-excel":
        if raw[:4] == OLE_COMPOUND_MAGIC:
            return "xls"
        if "openxml" in ctype:
            return "xlsx"
        return "xls"
    if raw[:2] == b"PK":
        return "xlsx"
    if raw[:4] == OLE_COMPOUND_MAGIC:
        return "xls"
    if ext in EXCEL_XLSX_EXTENSIONS:
        return "xlsx"
    if ext in CSV_EXTENSIONS:
        return "csv"
    raise TabularParseError(
        "Unsupported file type. Upload a .csv, .xlsx, or .xls file."
    )


def _decode_csv_bytes(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise TabularParseError("Could not decode CSV file (unsupported encoding).")


def _table_bounds(row: list[str]) -> tuple[int, int] | None:
    """Return the non-empty portion of a potential table header."""
    non_empty = [index for index, cell in enumerate(row) if cell]
    if not non_empty:
        return None
    return non_empty[0], non_empty[-1] + 1


def _header_score(
    rows_raw: list[list[str]],
    *,
    header_index: int,
    start_col: int,
    end_col: int,
) -> int:
    """
    Prefer dense label rows over title/preamble rows.

    This intentionally stays deterministic: it handles common spreadsheet
    exports with introductory text without sending uploaded list data to an AI.
    """
    headers = rows_raw[header_index][start_col:end_col]
    width = len(headers)
    label_cells = sum(
        1 for value in headers if value and not any(char.isdigit() for char in value)
    )
    supporting_rows = 0
    min_data_cells = max(1, (width + 1) // 2)
    for row in rows_raw[header_index + 1 : header_index + 4]:
        cells = row[start_col:end_col]
        if sum(1 for cell in cells if cell) >= min_data_cells:
            supporting_rows += 1
    return width * 100 + label_cells * 10 + supporting_rows


def _select_table(rows_raw: list[list[str]]) -> tuple[list[str], list[list[str]]]:
    """
    Select a likely header row from the opening section of an uploaded table.

    Leading/trailing blank cells are outside the table; a blank cell between
    two headers remains invalid because it makes a stable column mapping
    ambiguous.
    """
    if len(rows_raw) < 2:
        raise TabularParseError("File has headers but no data rows.")

    best: tuple[int, int, int, int] | None = None
    max_header_index = min(len(rows_raw) - 1, _HEADER_SCAN_ROWS)
    for header_index in range(max_header_index):
        bounds = _table_bounds(rows_raw[header_index])
        if bounds is None:
            continue
        start_col, end_col = bounds
        headers = rows_raw[header_index][start_col:end_col]
        if any(not header for header in headers):
            continue
        if len({header.casefold() for header in headers}) != len(headers):
            continue
        label_cells = sum(
            1
            for header in headers
            if not any(character.isdigit() for character in header)
        )
        if len(headers) > 1 and label_cells * 2 < len(headers):
            continue
        score = _header_score(
            rows_raw,
            header_index=header_index,
            start_col=start_col,
            end_col=end_col,
        )
        candidate = (score, header_index, start_col, end_col)
        if best is None or candidate[0] > best[0]:
            best = candidate

    if best is None:
        raise TabularParseError(
            "Could not detect a header row with unique, non-empty column names."
        )

    _, header_index, start_col, end_col = best
    return (
        rows_raw[header_index][start_col:end_col],
        [row[start_col:end_col] for row in rows_raw[header_index + 1 :]],
    )


def _parse_csv(raw: bytes) -> dict:
    text = _decode_csv_bytes(raw)
    reader = csv.reader(io.StringIO(text))
    rows_raw: list[list[str]] = []
    for row in reader:
        cells = [_normalize_cell(c) for c in row]
        if not any(cells):
            continue
        rows_raw.append(cells)
        if len(rows_raw) >= _MAX_COLLECTED_ROWS:
            break
    if not rows_raw:
        raise TabularParseError("CSV file has no data rows.")
    headers, data_rows = _select_table(rows_raw)
    return _table_from_rows(headers, data_rows, sheet_name=None)


def _parse_xlsx(raw: bytes) -> dict:
    from io import BytesIO

    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    try:
        sheet_name = workbook.sheetnames[0]
        sheet = workbook[sheet_name]
        rows_raw: list[list[str]] = []
        for row in sheet.iter_rows(values_only=True):
            cells = [_normalize_cell(c) for c in row[:_MAX_COLS]]
            if not any(cells):
                continue
            rows_raw.append(cells)
            if len(rows_raw) >= _MAX_COLLECTED_ROWS:
                break
    finally:
        workbook.close()
    if not rows_raw:
        raise TabularParseError("Excel file has no readable rows.")
    headers, data_rows = _select_table(rows_raw)
    return _table_from_rows(headers, data_rows, sheet_name=sheet_name)


def _parse_xls(raw: bytes) -> dict:
    try:
        import xlrd
    except ModuleNotFoundError as exc:
        raise TabularParseError(
            "Excel .xls support is not installed on this server."
        ) from exc

    try:
        workbook = xlrd.open_workbook(file_contents=raw)
    except Exception as exc:
        raise TabularParseError(f"Could not read Excel file: {exc}") from exc

    sheet = workbook.sheet_by_index(0)
    rows_raw: list[list[str]] = []
    max_col = min(sheet.ncols, _MAX_COLS)
    for rx in range(sheet.nrows):
        cells = []
        for cx in range(max_col):
            cell = sheet.cell(rx, cx)
            cells.append(_normalize_cell(cell.value))
        if not any(cells):
            continue
        rows_raw.append(cells)
        if len(rows_raw) >= _MAX_COLLECTED_ROWS:
            break
    if not rows_raw:
        raise TabularParseError("Excel file has no readable rows.")
    headers, data_rows = _select_table(rows_raw)
    return _table_from_rows(headers, data_rows, sheet_name=sheet.name)


def _table_from_rows(
    headers: list[str],
    data_rows: list[list[str]],
    *,
    sheet_name: str | None,
) -> dict:
    normalized_headers = [(h or "").strip() for h in headers]
    if not normalized_headers or not any(normalized_headers):
        raise TabularParseError("Header row is missing or empty.")
    seen: set[str] = set()
    for header in normalized_headers:
        if not header:
            raise TabularParseError("Header row contains an empty column name.")
        key = header.casefold()
        if key in seen:
            raise TabularParseError(f"Duplicate column name: {header}")
        seen.add(key)

    rows: list[dict[str, str]] = []
    for raw_row in data_rows[:_MAX_DATA_ROWS]:
        padded = list(raw_row) + [""] * (len(normalized_headers) - len(raw_row))
        row_dict = {
            normalized_headers[i]: _normalize_cell(padded[i])
            for i in range(len(normalized_headers))
        }
        if any(row_dict.values()):
            rows.append(row_dict)

    if not rows:
        raise TabularParseError("File has headers but no data rows.")

    return {
        "headers": normalized_headers,
        "rows": rows,
        "sheet_name": sheet_name,
    }


def parse_tabular_upload(
    raw: bytes,
    *,
    filename: str = "",
    content_type: str = "",
) -> dict:
    if not raw:
        raise TabularParseError("Uploaded file is empty.")
    fmt = _detect_format(raw, filename, content_type)
    if fmt == "csv":
        return _parse_csv(raw)
    if fmt == "xlsx":
        return _parse_xlsx(raw)
    if fmt == "xls":
        return _parse_xls(raw)
    raise TabularParseError("Unsupported file type.")


def build_config_from_table(headers: list[str], rows: list[dict[str, str]]) -> dict:
    columns: list[ColumnConfig] = []
    for header in headers:
        examples: list[str] = []
        empty_seen = False
        for row in rows:
            value = row.get(header, "")
            if not value:
                empty_seen = True
                continue
            if value not in examples:
                examples.append(value)
            if len(examples) >= 5:
                break
        columns.append(
            ColumnConfig(
                name=header,
                examples=examples,
                can_be_empty=empty_seen,
            )
        )
    config = OrganizationListConfig(columns=columns)
    return validate_list_config_for_storage(config.model_dump())
